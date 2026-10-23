from __future__ import annotations

import hashlib
import hmac
import json
import os
import secrets
import sqlite3
import time
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import httpx
from dotenv import load_dotenv
from docx import Document
from pypdf import PdfReader
from fastapi import Depends, FastAPI, File, Header, HTTPException, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, StreamingResponse
from pydantic import BaseModel, Field


ROOT = Path(__file__).resolve().parent
load_dotenv(ROOT / ".env")
DATA_DIR = Path(os.getenv("MINO_DATA_DIR", ROOT / "data"))
DB_PATH = Path(os.getenv("MINO_DATABASE_PATH", DATA_DIR / "mino.db"))
UPLOAD_DIR = Path(os.getenv("MINO_UPLOAD_DIR", DATA_DIR / "uploads"))
SECRET = os.getenv("MINO_SECRET_KEY", "development-only-change-this-secret")
ACCESS_SECONDS = int(os.getenv("MINO_ACCESS_TOKEN_MINUTES", "60")) * 60
REFRESH_SECONDS = int(os.getenv("MINO_REFRESH_TOKEN_DAYS", "30")) * 86400
MAX_UPLOAD = int(os.getenv("MINO_MAX_UPLOAD_MB", "20")) * 1024 * 1024
DATA_DIR.mkdir(parents=True, exist_ok=True)
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)


def now() -> str:
    return datetime.now(timezone.utc).isoformat()


def db() -> sqlite3.Connection:
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def init_db() -> None:
    with db() as c:
        c.executescript("""
        CREATE TABLE IF NOT EXISTS users(id TEXT PRIMARY KEY, username TEXT UNIQUE NOT NULL, name TEXT NOT NULL, role TEXT NOT NULL, password_hash TEXT NOT NULL, created_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS tokens(token_hash TEXT PRIMARY KEY, user_id TEXT NOT NULL REFERENCES users(id) ON DELETE CASCADE, kind TEXT NOT NULL, expires INTEGER NOT NULL);
        CREATE TABLE IF NOT EXISTS profiles(user_id TEXT PRIMARY KEY REFERENCES users(id) ON DELETE CASCADE, body TEXT NOT NULL DEFAULT '{}');
        CREATE TABLE IF NOT EXISTS learning(user_id TEXT PRIMARY KEY REFERENCES users(id) ON DELETE CASCADE, body TEXT NOT NULL DEFAULT '{}');
        CREATE TABLE IF NOT EXISTS sessions(id TEXT PRIMARY KEY, user_id TEXT NOT NULL REFERENCES users(id) ON DELETE CASCADE, title TEXT NOT NULL, parent_id TEXT, branch_message_id TEXT, created_at TEXT NOT NULL, updated_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS messages(id TEXT PRIMARY KEY, session_id TEXT NOT NULL REFERENCES sessions(id) ON DELETE CASCADE, role TEXT NOT NULL, body TEXT NOT NULL, created_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS attachments(id TEXT PRIMARY KEY, user_id TEXT NOT NULL REFERENCES users(id) ON DELETE CASCADE, name TEXT NOT NULL, mime_type TEXT NOT NULL, size INTEGER NOT NULL, path TEXT NOT NULL, created_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS feedback(id TEXT PRIMARY KEY, user_id TEXT NOT NULL REFERENCES users(id) ON DELETE CASCADE, body TEXT NOT NULL, created_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS interviews(id TEXT PRIMARY KEY, user_id TEXT NOT NULL REFERENCES users(id) ON DELETE CASCADE, config TEXT NOT NULL, status TEXT NOT NULL, report TEXT NOT NULL DEFAULT '{}', created_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS interview_answers(id TEXT PRIMARY KEY, session_id TEXT NOT NULL REFERENCES interviews(id) ON DELETE CASCADE, body TEXT NOT NULL, created_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS reviews(id TEXT PRIMARY KEY, body TEXT NOT NULL, status TEXT NOT NULL DEFAULT 'pending', decision TEXT, created_at TEXT NOT NULL);
        """)
        username = os.getenv("MINO_BOOTSTRAP_USERNAME", "demo")
        if not c.execute("SELECT id FROM users WHERE username=?", (username,)).fetchone():
            uid = "usr_" + uuid.uuid4().hex[:16]
            password = os.getenv("MINO_BOOTSTRAP_PASSWORD", "change-this-password")
            c.execute("INSERT INTO users VALUES(?,?,?,?,?,?)", (uid, username, username, os.getenv("MINO_BOOTSTRAP_ROLE", "user"), hash_password(password), now()))
            c.execute("INSERT INTO profiles(user_id,body) VALUES(?, '{}')", (uid,))
            c.execute("INSERT INTO learning(user_id,body) VALUES(?, '{}')", (uid,))


def hash_password(password: str, salt: bytes | None = None) -> str:
    salt = salt or secrets.token_bytes(16)
    key = hashlib.pbkdf2_hmac("sha256", password.encode(), salt, 310_000)
    return salt.hex() + ":" + key.hex()


def verify_password(password: str, stored: str) -> bool:
    try:
        salt_hex, key_hex = stored.split(":", 1)
        return hmac.compare_digest(hash_password(password, bytes.fromhex(salt_hex)), salt_hex + ":" + key_hex)
    except ValueError:
        return False


def token(kind: str, user_id: str, ttl: int) -> str:
    raw = secrets.token_urlsafe(40)
    with db() as c:
        c.execute("INSERT INTO tokens VALUES(?,?,?,?)", (hashlib.sha256(raw.encode()).hexdigest(), user_id, kind, int(time.time()) + ttl))
    return raw


def public_user(row: sqlite3.Row) -> dict[str, str]:
    return {"id": row["id"], "name": row["name"], "role": row["role"]}


def current_user(authorization: str | None = Header(default=None)) -> sqlite3.Row:
    if not authorization or not authorization.startswith("Bearer "):
        raise HTTPException(401, "需要登录")
    raw = authorization[7:]
    with db() as c:
        row = c.execute("SELECT users.* FROM tokens JOIN users ON users.id=tokens.user_id WHERE token_hash=? AND kind='access' AND expires>?", (hashlib.sha256(raw.encode()).hexdigest(), int(time.time()))).fetchone()
    if not row:
        raise HTTPException(401, "登录已失效")
    return row


def admin_user(user: sqlite3.Row = Depends(current_user)) -> sqlite3.Row:
    if user["role"] != "admin":
        raise HTTPException(403, "需要管理员权限")
    return user


def extract_file_text(path: str, mime_type: str) -> str:
    """Extract text for model context; scanned/image-only PDFs need OCR later."""
    try:
        if mime_type == "application/pdf":
            reader = PdfReader(path)
            return "\n\n".join((page.extract_text() or "") for page in reader.pages).strip()
        if mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            doc = Document(path)
            return "\n".join(p.text for p in doc.paragraphs).strip()
        return Path(path).read_text(encoding="utf-8", errors="replace").strip()
    except Exception:
        return ""


def attachment_context(user_id: str, attachment_ids: list[str]) -> str:
    if not attachment_ids:
        return ""
    chunks = []
    with db() as c:
        placeholders = ",".join("?" for _ in attachment_ids)
        rows = c.execute(f"SELECT id,name,mime_type,path FROM attachments WHERE user_id=? AND id IN ({placeholders})", [user_id, *attachment_ids]).fetchall()
    for row in rows:
        text = extract_file_text(row["path"], row["mime_type"])
        if text:
            chunks.append(f"[附件：{row['name']}]\n{text[:30000]}")
    return "\n\n".join(chunks)


class Login(BaseModel):
    username: str
    password: str


class Refresh(BaseModel):
    refreshToken: str


class JsonBody(BaseModel):
    model_config = {"extra": "allow"}

    def data(self) -> dict[str, Any]:
        return self.model_dump()


def body_dict(model: JsonBody) -> dict[str, Any]:
    return model.model_dump(exclude_unset=False)


app = FastAPI(title="MINO AGENT API", version="0.1.0", description="MINO 求职智能体后端 API")
origins = [x.strip() for x in os.getenv("MINO_CORS_ORIGINS", "http://localhost:5173,http://127.0.0.1:5173").split(",") if x.strip()]
app.add_middleware(CORSMiddleware, allow_origins=origins, allow_credentials=True, allow_methods=["*"], allow_headers=["*"])


@app.middleware("http")
async def request_id_middleware(request: Request, call_next):
    request_id = request.headers.get("x-request-id", str(uuid.uuid4()))
    try:
        response = await call_next(request)
    except Exception:
        response = JSONResponse(status_code=500, content={"error": {"code": "INTERNAL_ERROR", "message": "服务暂时不可用", "details": []}, "requestId": request_id})
    response.headers["X-Request-ID"] = request_id
    return response


@app.exception_handler(HTTPException)
async def http_error(request: Request, exc: HTTPException):
    codes = {401: "UNAUTHENTICATED", 403: "FORBIDDEN", 404: "NOT_FOUND", 409: "CONFLICT", 413: "FILE_TOO_LARGE"}
    return JSONResponse(status_code=exc.status_code, content={"error": {"code": codes.get(exc.status_code, "REQUEST_ERROR"), "message": str(exc.detail), "details": []}, "requestId": request.headers.get("x-request-id")})


@app.on_event("startup")
def startup() -> None:
    init_db()


@app.get("/health")
def health():
    return {"status": "ok", "time": now(), "storage": "sqlite", "modelConfigured": bool(os.getenv("MINO_LLM_BASE_URL") and os.getenv("MINO_LLM_API_KEY") and os.getenv("MINO_LLM_MODEL")), "model": os.getenv("MINO_LLM_MODEL") or None}


@app.post("/api/auth/login")
def login(payload: Login):
    with db() as c:
        row = c.execute("SELECT * FROM users WHERE username=?", (payload.username,)).fetchone()
    if not row or not verify_password(payload.password, row["password_hash"]):
        raise HTTPException(401, "用户名或密码错误")
    return {"accessToken": token("access", row["id"], ACCESS_SECONDS), "refreshToken": token("refresh", row["id"], REFRESH_SECONDS), "expiresIn": ACCESS_SECONDS, "user": public_user(row)}


@app.post("/api/auth/refresh")
def refresh(payload: Refresh):
    digest = hashlib.sha256(payload.refreshToken.encode()).hexdigest()
    with db() as c:
        row = c.execute("SELECT users.* FROM tokens JOIN users ON users.id=tokens.user_id WHERE token_hash=? AND kind='refresh' AND expires>?", (digest, int(time.time()))).fetchone()
        if not row:
            raise HTTPException(401, "刷新令牌无效")
        c.execute("DELETE FROM tokens WHERE token_hash=?", (digest,))
    return {"accessToken": token("access", row["id"], ACCESS_SECONDS), "refreshToken": token("refresh", row["id"], REFRESH_SECONDS), "expiresIn": ACCESS_SECONDS}


@app.post("/api/auth/logout")
def logout(user: sqlite3.Row = Depends(current_user), authorization: str = Header(...)):
    with db() as c:
        c.execute("DELETE FROM tokens WHERE user_id=?", (user["id"],))
    return {"ok": True}


@app.get("/api/users/me")
def me(user: sqlite3.Row = Depends(current_user)):
    return public_user(user)


@app.get("/api/users/me/profile")
def get_profile(user: sqlite3.Row = Depends(current_user)):
    with db() as c:
        row = c.execute("SELECT body FROM profiles WHERE user_id=?", (user["id"],)).fetchone()
    return {"profile": json.loads(row["body"] or "{}")}


@app.put("/api/users/me/profile")
def put_profile(payload: JsonBody, user: sqlite3.Row = Depends(current_user)):
    data = body_dict(payload)
    with db() as c:
        c.execute("INSERT INTO profiles(user_id,body) VALUES(?,?) ON CONFLICT(user_id) DO UPDATE SET body=excluded.body", (user["id"], json.dumps(data, ensure_ascii=False)))
    return {"profile": data}


@app.get("/api/users/me/learning")
def get_learning(user: sqlite3.Row = Depends(current_user)):
    with db() as c:
        row = c.execute("SELECT body FROM learning WHERE user_id=?", (user["id"],)).fetchone()
    return json.loads(row["body"] or "{}")


@app.put("/api/users/me/learning")
def put_learning(payload: JsonBody, user: sqlite3.Row = Depends(current_user)):
    data = body_dict(payload)
    with db() as c:
        c.execute("INSERT INTO learning(user_id,body) VALUES(?,?) ON CONFLICT(user_id) DO UPDATE SET body=excluded.body", (user["id"], json.dumps(data, ensure_ascii=False)))
    return {"state": data}


@app.get("/api/usage/me")
def usage(user: sqlite3.Row = Depends(current_user)):
    with db() as c:
        count = c.execute("SELECT COUNT(*) n FROM messages JOIN sessions ON sessions.id=messages.session_id WHERE sessions.user_id=? AND date(messages.created_at)=date('now')", (user["id"],)).fetchone()["n"]
    return {"today": count, "limit": None}


@app.delete("/api/users/me/data")
def delete_user_data(user: sqlite3.Row = Depends(current_user)):
    with db() as c:
        c.execute("DELETE FROM sessions WHERE user_id=?", (user["id"],))
        c.execute("DELETE FROM profiles WHERE user_id=?", (user["id"],))
        c.execute("DELETE FROM learning WHERE user_id=?", (user["id"],))
        c.execute("DELETE FROM feedback WHERE user_id=?", (user["id"],))
        c.execute("DELETE FROM interviews WHERE user_id=?", (user["id"],))
        c.execute("DELETE FROM attachments WHERE user_id=?", (user["id"],))
    return {"deleted": True}


@app.get("/api/knowledge/colleges")
def colleges(): return {"items": []}
@app.get("/api/knowledge/colleges/{college_id}")
def college(college_id: str): raise HTTPException(404, "院校数据尚未导入")
@app.get("/api/knowledge/majors/{major_id}")
def major(major_id: str): raise HTTPException(404, "专业数据尚未导入")
@app.get("/api/knowledge/textbooks")
def textbooks(): return {"items": []}
@app.get("/api/guides")
def guides(): return {"items": []}
@app.get("/api/guides/{guide_id}")
def guide(guide_id: str): raise HTTPException(404, "求职指南尚未导入")
@app.get("/api/guides/{guide_id}/quiz")
def quiz(guide_id: str): return {"items": []}


@app.get("/api/chat/sessions")
def list_sessions(user: sqlite3.Row = Depends(current_user)):
    with db() as c:
        rows = c.execute("SELECT * FROM sessions WHERE user_id=? ORDER BY updated_at DESC", (user["id"],)).fetchall()
    return {"items": [dict(r) for r in rows]}


@app.post("/api/chat/sessions")
def create_session(payload: JsonBody, user: sqlite3.Row = Depends(current_user)):
    data = body_dict(payload); sid = "conv_" + uuid.uuid4().hex[:16]
    with db() as c:
        if data.get("parentSessionId"):
            parent = c.execute("SELECT id FROM sessions WHERE id=? AND user_id=?", (data["parentSessionId"], user["id"])).fetchone()
            if not parent: raise HTTPException(404, "源会话不存在")
        c.execute("INSERT INTO sessions VALUES(?,?,?,?,?,?,?)", (sid, user["id"], data.get("title") or "新对话", data.get("parentSessionId"), data.get("branchMessageId"), now(), now()))
    return {"id": sid, "title": data.get("title") or "新对话", "parentSessionId": data.get("parentSessionId"), "branchMessageId": data.get("branchMessageId")}


@app.patch("/api/chat/sessions/{session_id}")
def rename_session(session_id: str, payload: JsonBody, user: sqlite3.Row = Depends(current_user)):
    title = str(payload.model_dump().get("title", "")).strip()
    if not title: raise HTTPException(422, "标题不能为空")
    with db() as c:
        cur = c.execute("UPDATE sessions SET title=?,updated_at=? WHERE id=? AND user_id=?", (title[:120], now(), session_id, user["id"]))
        if not cur.rowcount: raise HTTPException(404, "会话不存在")
    return {"id": session_id, "title": title[:120]}


@app.delete("/api/chat/sessions/{session_id}")
def delete_session(session_id: str, user: sqlite3.Row = Depends(current_user)):
    with db() as c:
        cur = c.execute("DELETE FROM sessions WHERE id=? AND user_id=?", (session_id, user["id"]))
        if not cur.rowcount: raise HTTPException(404, "会话不存在")
    return {"id": session_id, "deleted": True}


@app.get("/api/chat/sessions/{session_id}/messages")
def list_session_messages(session_id: str, user: sqlite3.Row = Depends(current_user)):
    with db() as c:
        session = c.execute("SELECT id FROM sessions WHERE id=? AND user_id=?", (session_id, user["id"])).fetchone()
        if not session: raise HTTPException(404, "会话不存在")
        rows = c.execute("SELECT * FROM messages WHERE session_id=? ORDER BY created_at", (session_id,)).fetchall()
    items = []
    for row in rows:
        body = json.loads(row["body"])
        if row["role"] == "user":
            content = body.get("text", "")
            references = ([{"kind": "knowledge", "title": body["context"].get("title", "知识库")} ] if body.get("context") else [])
            attachments = body.get("attachments", [])
        else:
            content = body.get("text", "")
            references, attachments = [], []
        items.append({"id": row["id"], "role": row["role"], "content": content, "references": references, "attachments": attachments, "createdAt": row["created_at"]})
    return {"items": items}


async def model_stream(text: str, context: dict[str, Any] | None, history: list[dict[str, str]]):
    base = os.getenv("MINO_LLM_BASE_URL", "").rstrip("/")
    key, model = os.getenv("MINO_LLM_API_KEY", ""), os.getenv("MINO_LLM_MODEL", "")
    if base and key and model:
        async with httpx.AsyncClient(timeout=90) as client:
            messages = [{"role": "system", "content": "你是 MINO AGENT 求职智能体。给出具体、可执行、诚实的求职建议。"}, *history, {"role": "user", "content": (json.dumps(context, ensure_ascii=False) + "\n" if context else "") + text}]
            async with client.stream("POST", base + "/chat/completions", headers={"Authorization": f"Bearer {key}"}, json={"model": model, "stream": True, "messages": messages}) as res:
                if res.status_code >= 400:
                    await res.aread()
                res.raise_for_status()
                async for line in res.aiter_lines():
                    if not line.startswith("data:"):
                        continue
                    value = line[5:].strip()
                    if value == "[DONE]":
                        break
                    try:
                        token_data = json.loads(value)["choices"][0].get("delta", {}).get("content")
                    except (ValueError, KeyError, IndexError, TypeError):
                        continue
                    if token_data:
                        yield token_data
        return
    topic = text[:80] or "你的求职问题"
    fallback = f"## 关于「{topic}」\n\n这是本地开发模式生成的示例回复，尚未连接语言模型。\n\n建议先补充目标岗位、相关经历和可量化结果，我可以据此整理岗位匹配点、STAR 案例和下一步行动。"
    for offset in range(0, len(fallback), 48):
        yield fallback[offset:offset + 48]


def sse(event: str, data: dict[str, Any]) -> str:
    return f"event: {event}\ndata: {json.dumps(data, ensure_ascii=False)}\n\n"


@app.post("/api/chat")
async def chat(payload: JsonBody, request: Request, user: sqlite3.Row = Depends(current_user)):
    data = body_dict(payload); text = str(data.get("text", "")).strip()
    attachment_ids = [str(item) for item in data.get("attachmentIds", []) if item]
    attachment_text = attachment_context(user["id"], attachment_ids)
    if not text and not attachment_text: raise HTTPException(422, "消息不能为空或附件无法提取文字")
    model_text = text
    if attachment_text:
        model_text += ("\n\n" if model_text else "") + "请基于以下附件内容回答：\n" + attachment_text
    session_id = data.get("conversationId")
    with db() as c:
        if session_id:
            row = c.execute("SELECT id FROM sessions WHERE id=? AND user_id=?", (session_id, user["id"])).fetchone()
            if not row: raise HTTPException(404, "会话不存在")
        else:
            session_id = "conv_" + uuid.uuid4().hex[:16]
            c.execute("INSERT INTO sessions VALUES(?,?,?,?,?,?,?)", (session_id, user["id"], text[:40], None, None, now(), now()))
        previous_rows = c.execute("SELECT role,body FROM messages WHERE session_id=? ORDER BY created_at DESC LIMIT 20", (session_id,)).fetchall()
        history = []
        for previous in reversed(previous_rows):
            previous_body = json.loads(previous["body"])
            content = previous_body.get("text", "") if previous["role"] == "user" else previous_body.get("text", "")
            if content:
                history.append({"role": previous["role"], "content": content})
        c.execute("INSERT INTO messages VALUES(?,?,?,?,?)", ("msg_" + uuid.uuid4().hex[:16], session_id, "user", json.dumps(data, ensure_ascii=False), now()))
    async def events():
        yield sse("activity", {"status": "running", "label": "正在整理求职建议"})
        try:
            parts = []
            async for part in model_stream(model_text, data.get("context"), history):
                if await request.is_disconnected(): return
                parts.append(part)
                yield sse("delta", {"delta": part})
            answer = "".join(parts)
            mid = "msg_" + uuid.uuid4().hex[:16]
            with db() as c:
                c.execute("INSERT INTO messages VALUES(?,?,?,?,?)", (mid, session_id, "assistant", json.dumps({"text": answer, "mode": data.get("mode"), "tool": data.get("tool")}, ensure_ascii=False), now()))
                c.execute("UPDATE sessions SET updated_at=? WHERE id=?", (now(), session_id))
            yield sse("done", {"messageId": mid, "usage": None, "performance": None, "citations": []})
        except httpx.HTTPStatusError as exc:
            if exc.response.status_code in (401, 403):
                yield sse("error", {"code": "MODEL_AUTH_FAILED", "message": "模型服务鉴权失败，请检查服务端 API Key 与所属地域", "retryable": False})
            elif exc.response.status_code == 400:
                try:
                    provider_error = exc.response.json().get("error", {})
                except (ValueError, AttributeError):
                    provider_error = {}
                if str(provider_error.get("code")) == "2030":
                    yield sse("error", {"code": "MODEL_MIGRATION_REQUIRED", "message": "当前混元凭据对应的旧版模型服务已迁移，请按腾讯云 TokenHub 迁移说明更新服务地址和 API 凭据", "retryable": False})
                else:
                    yield sse("error", {"code": "MODEL_BAD_REQUEST", "message": "模型服务拒绝了请求，请检查模型名称和服务配置", "retryable": False})
            else:
                yield sse("error", {"code": "MODEL_UNAVAILABLE", "message": f"模型服务暂时不可用（HTTP {exc.response.status_code}）", "retryable": exc.response.status_code >= 500})
        except Exception as exc:
            yield sse("error", {"code": "MODEL_UNAVAILABLE", "message": "暂时无法连接模型服务，请检查网络和模型配置", "retryable": True, "diagnostic": type(exc).__name__})
    return StreamingResponse(events(), media_type="text/event-stream", headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


@app.post("/api/chat/attachments")
async def upload(files: list[UploadFile] = File(...), user: sqlite3.Row = Depends(current_user)):
    accepted = []
    allowed = {"application/pdf", "text/plain", "text/markdown", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "application/msword"}
    for file in files:
        contents = await file.read(MAX_UPLOAD + 1)
        if len(contents) > MAX_UPLOAD: raise HTTPException(413, "单个文件超过大小限制")
        mime = file.content_type or "application/octet-stream"
        if mime not in allowed: raise HTTPException(422, f"不支持的文件类型: {mime}")
        aid = "file_" + uuid.uuid4().hex[:16]
        safe_name = Path(file.filename or "upload").name
        target = UPLOAD_DIR / (aid + Path(safe_name).suffix.lower()[:12])
        target.write_bytes(contents)
        with db() as c: c.execute("INSERT INTO attachments VALUES(?,?,?,?,?,?,?)", (aid, user["id"], safe_name, mime, len(contents), str(target), now()))
        accepted.append({"id": aid, "name": safe_name, "mimeType": mime, "size": len(contents), "status": "ready"})
    return {"files": accepted}


@app.post("/api/chat/feedback")
def chat_feedback(payload: JsonBody, user: sqlite3.Row = Depends(current_user)):
    data = body_dict(payload)
    if data.get("rating") not in {"helpful", "unhelpful"}: raise HTTPException(422, "rating 值无效")
    fid = "fb_" + uuid.uuid4().hex[:16]
    with db() as c: c.execute("INSERT INTO feedback VALUES(?,?,?,?)", (fid, user["id"], json.dumps(data, ensure_ascii=False), now()))
    return {**data, "id": fid, "accepted": True}


@app.get("/api/interview/questions")
def interview_questions(position: str = "通用", difficulty: str = "初级"):
    return {"position": position, "difficulty": difficulty, "items": []}


def feedback_for(answer: str) -> dict[str, Any]:
    words = len(answer.strip())
    return {"score": 3 if words else 1, "dimensions": [{"name": "表达完整度", "stars": 3 if words else 1, "note": "请结合具体背景、行动和结果展开。"}], "strengths": [], "suggestions": ["补充可量化结果", "使用 STAR 结构组织回答"], "starExample": "", "relatedGuideIds": []}


@app.post("/api/interview/sessions")
def create_interview(payload: JsonBody, user: sqlite3.Row = Depends(current_user)):
    data = body_dict(payload); sid = "int_" + uuid.uuid4().hex[:16]
    with db() as c: c.execute("INSERT INTO interviews VALUES(?,?,?,?,?,?)", (sid, user["id"], json.dumps(data, ensure_ascii=False), "active", "{}", now()))
    return {"id": sid, **data, "status": "active"}


@app.post("/api/interview/sessions/{session_id}/answers")
def interview_answer(session_id: str, payload: JsonBody, user: sqlite3.Row = Depends(current_user)):
    data = body_dict(payload)
    with db() as c:
        row = c.execute("SELECT config FROM interviews WHERE id=? AND user_id=? AND status='active'", (session_id, user["id"])).fetchone()
        if not row: raise HTTPException(404, "面试会话不存在或已结束")
        config = json.loads(row["config"])
        c.execute("INSERT INTO interview_answers VALUES(?,?,?,?)", ("ans_" + uuid.uuid4().hex[:16], session_id, json.dumps(data, ensure_ascii=False), now()))
    return {"sessionId": session_id, **data, "feedback": feedback_for(data.get("answer", ""))}


@app.post("/api/interview/sessions/{session_id}/finish")
def finish_interview(session_id: str, user: sqlite3.Row = Depends(current_user)):
    with db() as c:
        row = c.execute("SELECT id FROM interviews WHERE id=? AND user_id=?", (session_id, user["id"])).fetchone()
        if not row: raise HTTPException(404, "面试会话不存在")
        answers = [json.loads(x["body"]) for x in c.execute("SELECT body FROM interview_answers WHERE session_id=?", (session_id,)).fetchall()]
        report = {"answerCount": len(answers), "summary": "面试已结束，建议结合逐题反馈复盘。", "answers": answers}
        c.execute("UPDATE interviews SET status='completed',report=? WHERE id=?", (json.dumps(report, ensure_ascii=False), session_id))
    return {"sessionId": session_id, "status": "completed", "report": report}


@app.post("/api/interview/feedback")
def legacy_interview_feedback(payload: JsonBody, user: sqlite3.Row = Depends(current_user)):
    data = body_dict(payload)
    return feedback_for(data.get("answer", ""))


@app.post("/api/career/resume/analyze")
def resume_analyze(payload: JsonBody, user: sqlite3.Row = Depends(current_user)):
    data = body_dict(payload); text = data.get("text", "")
    return {"summary": "简历已接收。连接语言模型后可生成个性化分析。" if text or data.get("attachmentIds") else "请先提供简历内容。", "sections": [], "suggestions": ["补充量化成果", "按目标岗位调整关键词"] if text else [], "jobMatches": []}


@app.post("/api/career/jobs/match")
def job_match(payload: JsonBody, user: sqlite3.Row = Depends(current_user)):
    return {"items": [], "profile": body_dict(payload).get("profile", {})}


@app.post("/api/career/star/organize")
def star(payload: JsonBody, user: sqlite3.Row = Depends(current_user)):
    data = body_dict(payload); exp = data.get("experience", "")
    return {"situation": "", "task": "", "action": exp, "result": "", "suggestions": ["补充背景和目标", "描述个人行动", "增加可量化结果"]}


@app.get("/api/review/drafts")
def review_drafts(user: sqlite3.Row = Depends(admin_user)):
    with db() as c: rows = c.execute("SELECT * FROM reviews WHERE status='pending' ORDER BY created_at").fetchall()
    return {"items": [dict(r) for r in rows]}


@app.post("/api/review/drafts/{draft_id}/decision")
def review_decision(draft_id: str, payload: JsonBody, user: sqlite3.Row = Depends(admin_user)):
    data = body_dict(payload); decision = data.get("decision")
    if decision not in {"approved", "rejected"}: raise HTTPException(422, "decision 必须是 approved 或 rejected")
    with db() as c:
        cur = c.execute("UPDATE reviews SET status=?,decision=? WHERE id=? AND status='pending'", (decision, json.dumps(data, ensure_ascii=False), draft_id))
        if not cur.rowcount: raise HTTPException(404, "待审核内容不存在")
    return {"ok": True, "id": draft_id, "decision": decision}

