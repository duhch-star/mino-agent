"""Synthetic inputs only; no personal documents/configuration or network calls."""
import os
os.environ.setdefault('SECRET_KEY', 'synthetic-test-only-not-a-real-secret-12345678901234567890')
import tempfile
import struct
import unittest
from pathlib import Path
from unittest.mock import patch
from app.services import document_service as documents
from app.services.workspace_documents import read_workspace_document
from app.core.responses import ApiException


def pdf_bytes(text):
    stream = f'BT /F1 12 Tf 20 100 Td ({text}) Tj ET'.encode()
    objects = [b'<< /Type /Catalog /Pages 2 0 R >>', b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
        b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 200] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>',
        b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>', b'<< /Length ' + str(len(stream)).encode() + b' >>\nstream\n' + stream + b'\nendstream']
    data = bytearray(b'%PDF-1.4\n'); offsets = [0]
    for i, obj in enumerate(objects, 1):
        offsets.append(len(data)); data.extend(f'{i} 0 obj\n'.encode() + obj + b'\nendobj\n')
    start = len(data); data.extend(b'xref\n0 6\n0000000000 65535 f \n')
    for offset in offsets[1:]: data.extend(f'{offset:010} 00000 n \n'.encode())
    data.extend(f'trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{start}\n%%EOF'.encode())
    return data


class DocumentsTest(unittest.TestCase):
    def test_pdf_docx_content(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            (root / 'example.pdf').write_bytes(pdf_bytes('Synthetic PDF evidence'))
            self.assertIn('Synthetic PDF evidence', read_workspace_document(folder, 'example.pdf')['text'])
            doc = Document(); doc.add_paragraph('测试简历 synthetic DOCX')
            doc.add_table(rows=1, cols=1).cell(0, 0).text = 'table evidence'
            doc.save(root / 'example.docx')
            result = read_workspace_document(folder, str(root / 'example.docx'))
            self.assertIn('测试简历', result['text']); self.assertIn('table evidence', result['text'])

    def test_security_and_empty(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder); (root / 'blank.pdf').write_bytes(pdf_bytes(''))
            for requested in ['../outside.pdf', str(root.parent / 'outside.pdf'), 'x.txt', 'missing.pdf', 'blank.pdf']:
                with self.subTest(requested=requested), self.assertRaises(ApiException):
                    read_workspace_document(folder, requested)
            (root / 'invalid.doc').write_bytes(b'not Word')
            with self.assertRaises(ApiException): read_workspace_document(folder, 'invalid.doc')
            try:
                (root / 'escape.pdf').symlink_to(root.parent / 'outside.pdf')
            except OSError:
                pass  # Windows symlink creation may require a privilege.
            else:
                with self.assertRaises(ApiException): read_workspace_document(folder, 'escape.pdf')

    def test_bridge_verifies_identity_before_reading(self):
        from app.routers_fe.pi_web import document_read, DocumentRequest
        request = DocumentRequest(bridgeSecret='synthetic', token='synthetic', path='sample.pdf')
        with patch('app.routers_fe.pi_web.manager.verify', return_value={'workspace': 'verified-account-workspace'}) as verify, patch('app.services.workspace_documents.read_workspace_document', return_value={'text': 'parsed'}) as read:
            self.assertEqual(document_read(request, None), {'text': 'parsed'})
            verify.assert_called_once_with(None, token='synthetic', bridge_secret='synthetic')
            read.assert_called_once_with('verified-account-workspace', 'sample.pdf')
        with patch('app.routers_fe.pi_web.manager.verify', side_effect=ApiException(40100, 'invalid', status_code=401)), patch('app.services.workspace_documents.read_workspace_document') as read:
            with self.assertRaises(ApiException): document_read(request, None)
            read.assert_not_called()

    def test_truncation(self):
        with tempfile.TemporaryDirectory() as folder:
            (Path(folder) / 'long.pdf').touch()
            with patch('app.services.workspace_documents._extract_local', return_value='a' * 8001):
                result = read_workspace_document(folder, 'long.pdf')
                self.assertTrue(result['truncated']); self.assertEqual(len(result['text']), 8000)

    def test_word_piece_table_unicode_and_compressed(self):
        # A synthetic WordDocument stream and CLX exercise offsets, encodings and
        # main-story boundaries without scraping binary bytes as fake prose.
        word = bytearray(1024); word[:2] = b'\xec\xa5'
        struct.pack_into('<I', word, 76, 7)
        word[512:515] = b'ABC'; word[600:608] = '中文测试'.encode('utf-16-le')
        plc = struct.pack('<III', 0, 3, 7) + b'\0\0' + struct.pack('<I', 0x40000000 | (512 * 2)) + b'\0\0' + b'\0\0' + struct.pack('<I', 600) + b'\0\0'
        table = b'\x02' + struct.pack('<I', len(plc)) + plc
        struct.pack_into('<II', word, 418, 0, len(table))
        self.assertEqual(documents._doc_piece_text(word, table), 'ABC中文测试')
        import io
        for flag, table_name in [(0, '0Table'), (0x200, '1Table')]:
            struct.pack_into('<H', word, 10, flag)
            class SyntheticOle:
                def __enter__(self): return self
                def __exit__(self, *args): pass
                def openstream(self, name):
                    if name == 'WordDocument': return io.BytesIO(word)
                    self_test.assertEqual(name, table_name)
                    return io.BytesIO(table)
            self_test = self
            with patch('olefile.isOleFile', return_value=True), patch('olefile.OleFileIO', return_value=SyntheticOle()):
                self.assertEqual(documents._extract_doc_legacy('synthetic.doc'), 'ABC中文测试')
        struct.pack_into('<H', word, 10, 0x100)
        self.assertEqual(documents._doc_piece_text(word, table), '')
        self.assertEqual(documents._doc_piece_text(b'not a Word file', b''), '')


if __name__ == '__main__': unittest.main()
